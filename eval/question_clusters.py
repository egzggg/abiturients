"""Validate question uploads before atomically creating dashboard records."""
from __future__ import annotations

from contextlib import closing
from io import BytesIO
import json
from pathlib import Path
import sqlite3

MAX_QUESTIONS = 500
MAX_FILE_BYTES = 10 * 1024 * 1024


def parse_questions(content: bytes, filename: str) -> list[str]:
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError('Загрузите непустой файл размером до 10 МБ.')
    extension = Path(filename).suffix.lower()
    if extension == '.json':
        try:
            data = json.loads(content.decode('utf-8-sig'))
        except (ValueError, UnicodeError) as error:
            raise ValueError('Некорректный JSON. Используйте {"request": ["вопрос 1", "вопрос 2"]}.') from error
        if not isinstance(data, dict) or set(data) != {'request'} or not isinstance(data['request'], list):
            raise ValueError('JSON должен содержать только поле request с массивом вопросов.')
        values = data['request']
    elif extension in ('.xlsx', '.xls'):
        values = _excel_questions(content, extension)
    else:
        raise ValueError('Поддерживаются только файлы .json, .xlsx и .xls.')
    if not 1 <= len(values) <= MAX_QUESTIONS:
        raise ValueError(f'В файле должно быть от 1 до {MAX_QUESTIONS} вопросов.')
    questions = []
    for index, value in enumerate(values, 1):
        if not isinstance(value, str) or not value.strip() or len(value.strip()) > 20000:
            raise ValueError(f'Вопрос №{index}: нужен непустой текст длиной до 20 000 символов.')
        questions.append(value.strip())
    return questions


def _excel_questions(content: bytes, extension: str) -> list:
    workbook = None
    try:
        if extension == '.xlsx':
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
            rows = workbook.worksheets[0].iter_rows(values_only=True)
        else:
            import xlrd
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            sheet = workbook.sheet_by_index(0)
            rows = (sheet.row_values(index) for index in range(sheet.nrows))
        header = next(rows, ())
        matches = [index for index, value in enumerate(header) if isinstance(value, str) and value.strip() == 'request']
        if len(matches) != 1:
            raise ValueError('На первом листе в первой строке должен быть один столбец с названием request.')
        column = matches[0]
        questions = []
        for row_number, row in enumerate(rows, 2):
            value = row[column] if column < len(row) else None
            if value is None or isinstance(value, str) and not value.strip():
                continue
            if not isinstance(value, str) or value.startswith('='):
                raise ValueError(f'Строка {row_number}: в столбце request должен быть текст вопроса, а не число или формула.')
            questions.append(value)
            if len(questions) > MAX_QUESTIONS:
                raise ValueError(f'Можно загрузить до {MAX_QUESTIONS} вопросов за один раз.')
        return questions
    except ImportError as error:
        raise ValueError('На сервере нужны openpyxl и xlrd для загрузки Excel.') from error
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('Не удалось прочитать Excel. Проверьте формат и содержимое файла.') from error
    finally:
        if workbook is not None:
            if extension == '.xlsx':
                workbook.close()
            else:
                workbook.release_resources()


def import_questions(db_path: Path, content: bytes, filename: str, model: str) -> list[int]:
    from RAG_HP import initialize_sqlite_database
    from sqlite_answer_actions import generation_config
    questions = parse_questions(content, filename)
    if model not in generation_config()['models']:
        raise ValueError('Выберите модель из списка Yandex API.')
    initialize_sqlite_database(db_path)
    with closing(sqlite3.connect(db_path, timeout=5)) as connection, connection:
        # Validate the complete upload first; insertion is one transaction.
        ids = []
        for question in questions:
            cursor = connection.execute('INSERT INTO answers (request, llm_model) VALUES (?, ?)', (question, model))
            ids.append(cursor.lastrowid)
        return ids
